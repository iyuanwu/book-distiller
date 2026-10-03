"""Explicit real conversions of original fixtures. No private books."""
from pathlib import Path
import zipfile
import pytest
from book_distiller.core.ingest import IngestService
from book_distiller.core.parse import ParseService
from book_distiller.normalize.book import read_blocks

pytestmark = pytest.mark.docling_real


def test_real_pdf(isolated_storage):
    fixture = Path(__file__).parents[1] / "fixtures/document.pdf"
    manifest = IngestService(isolated_storage).ingest(fixture).manifest
    result = ParseService(isolated_storage).parse(manifest.book.slug)
    assert result.book.parser_metadata.parser == "docling"
    assert result.quality.total_pages == 2
    blocks = list(read_blocks(result.parsed_path / "normalized/blocks.jsonl"))
    assert any(block.type == "heading" for block in blocks)
    assert any(block.type == "paragraph" for block in blocks)
    assert {s.source_page_number for block in blocks for s in block.source_spans} == {1,2}
    assert all(s.source_page_index == s.source_page_number-1 for block in blocks for s in block.source_spans)
    assert [block.order for block in blocks] == list(range(1,len(blocks)+1))
    for name in ("raw/docling.json","raw/docling.md","raw/parse_metadata.json","normalized/book.json","normalized/blocks.jsonl","normalized/quality.json"):
        assert (result.parsed_path / name).is_file()
    assert ParseService(isolated_storage).parse(manifest.book.slug).already_parsed


@pytest.mark.parametrize("extension", [".md", ".markdown", ".docx", ".epub"])
def test_real_other_formats(isolated_storage, tmp_path, extension):
    fixture = Path(__file__).parents[1] / "fixtures/sample.md"
    source = tmp_path / ("original" + extension)
    if extension in (".md", ".markdown"):
        source.write_bytes(fixture.read_bytes())
    elif extension == ".docx":
        from docx import Document
        document = Document()
        document.add_heading("Original fixture", level=1)
        document.add_paragraph("A small notebook helps a reader compare two ideas.")
        document.save(source)
    else:
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("mimetype", "application/epub+zip")
            archive.writestr("META-INF/container.xml", '''<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>''')
            archive.writestr("content.opf", '''<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">original-fixture</dc:identifier><dc:title>Original fixture</dc:title><dc:language>en</dc:language></metadata><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>''')
            archive.writestr("chapter.xhtml", '''<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Original fixture</title></head><body><h1>Original fixture</h1><p>A small notebook helps a reader compare two ideas.</p></body></html>''')
    manifest = IngestService(isolated_storage).ingest(source).manifest
    result = ParseService(isolated_storage).parse(manifest.book.slug)
    assert result.book.block_count >= 2
    blocks = list(read_blocks(result.parsed_path / "normalized/blocks.jsonl"))
    assert any("notebook" in block.text for block in blocks)
    assert all(span.source_page_number is None for block in blocks for span in block.source_spans)
